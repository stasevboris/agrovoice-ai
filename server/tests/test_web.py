import io
from pathlib import Path
import re

from openpyxl import load_workbook
from PIL import Image
import pytest
from sqlalchemy import func, select

from app.models import Inspection, ReviewEvent, User
from app.security import issue_token, passwords
from app.web import boxes
from tests.test_api import client, database, migrate_test_database  # noqa: F401
from tests.test_sync import auth, other_auth, push, update  # noqa: F401


def csrf(html):
    return re.search(r'name="token" (?:type="hidden" )?value="([^"]+)"', html)[1]


def login(client, username="alice", password="correct-test-password"):
    page = client.get("/login")
    response = client.post("/login", data={"username": username, "password": password, "token": csrf(page.text)},
                           follow_redirects=False)
    assert response.status_code == 303
    return response


def sample_record(client, auth, note="Тестовая заметка", plot="Теплица №2"):
    image = io.BytesIO()
    Image.new("RGB", (20, 10), (70, 120, 50)).save(image, "PNG")
    item = update(image.getvalue())
    item["data"].update(note=note, plot=plot, lat=53.9, lon=27.56, created_at="2026-10-19T09:00:00+00:00")
    item["data"]["prediction"] = {"decoded_width": 20, "decoded_height": 10, "detections": [
        {"class_name": "potato__late_blight", "label_ru": "Картофель: фитофтороз", "xyxy_pixels": [1, 1, 8, 8]}]}
    response = push(client, auth, item)
    assert response.status_code == 200
    photo = item["data"]["photos"][0]
    assert client.put(f'/photos/{photo["id"]}', headers={**auth, "X-Content-SHA256": photo["sha256"]},
                      content=image.getvalue()).status_code == 200
    return item


def test_web_login_session_csrf_and_logout(client):
    assert client.get("/", follow_redirects=False).status_code == 303
    assert client.post("/login", data={"username": "alice", "password": "correct-test-password", "token": "bad"}).status_code == 403
    assert client.post("/login", data={"username": "alice", "password": "correct-test-password", "token": "чужой"}).status_code == 403
    page = client.get("/login")
    invalid = client.post("/login", data={"username": "alice", "password": "bad", "token": csrf(page.text)})
    assert invalid.status_code == 401 and "Не удалось войти" in invalid.text
    signed_in = login(client)
    page = client.get("/")
    assert "Осмотры растений" in page.text
    cookie = signed_in.headers["set-cookie"].lower()
    assert "httponly" in cookie and "samesite=strict" in cookie
    assert client.post("/logout", data={"token": "bad"}).status_code == 403
    assert client.post("/logout", data={"token": csrf(page.text)}, follow_redirects=False).status_code == 303
    assert client.get("/", follow_redirects=False).status_code == 303


def test_dashboard_filters_map_and_htmx_fragment(client, database, auth):
    item = sample_record(client, auth)
    login(client)
    page = client.get("/?crop=potato&disease=potato__late_blight&plot=Теплица%20№2&start=2026-10-19&end=2026-10-19")
    assert page.status_code == 200
    assert "Картофель: фитофтороз" in page.text
    assert "53.9" in page.text and 'data-points="' in page.text
    for query in ("crop=tomato", "disease=potato__healthy", "plot=Missing", "start=2026-10-20", "end=2026-10-18"):
        assert "Осмотров пока нет" in client.get("/?" + query).text
    fragment = client.get("/?crop=potato", headers={"HX-Request": "true"})
    assert "<html" not in fragment.text and item["id"] in fragment.text
    assert client.get("/?start=2026-10-20&end=2026-10-19").status_code == 422


def test_photo_boxes_and_escaping(client, database, auth):
    item = sample_record(client, auth, note='<script>alert("x")</script>')
    owner = database.scalar(select(User).where(User.username == "alice"))
    login(client)
    page = client.get(f'/inspections/{owner.id}/{item["id"]}')
    assert page.status_code == 200 and '<svg viewBox="0 0 20 10"' in page.text
    assert '<script>alert' not in page.text and '&lt;script&gt;' in page.text
    assert '<rect x="1" y="1" width="7" height="7">' in page.text
    photo = client.get(f'/preview/{owner.id}/{item["data"]["photos"][0]["id"]}')
    assert photo.status_code == 200 and photo.headers["content-type"] == "image/jpeg"
    assert Image.open(io.BytesIO(photo.content)).size == (20, 10)
    assert boxes({"prediction": {"decoded_width": 20, "decoded_height": 10,
                                "detections": [{"xyxy_pixels": [0, 0, 99, 99]}]}})[2] == []


def test_preview_limits_pixels_before_decoding(client, database, auth, monkeypatch):
    item = sample_record(client, auth)
    owner = database.scalar(select(User).where(User.username == "alice"))
    login(client)
    # Use a small fixture with a lowered limit; never allocate a decompression bomb.
    monkeypatch.setattr("app.web.MAX_PREVIEW_PIXELS", 100, raising=False)

    def forbidden_decode(*args, **kwargs):
        pytest.fail("Oversized preview must be rejected before decoding pixels")

    monkeypatch.setattr(Image.Image, "load", forbidden_decode)
    response = client.get(f'/preview/{owner.id}/{item["data"]["photos"][0]["id"]}')
    assert response.status_code == 415
    original = client.get(f'/photos/{item["data"]["photos"][0]["id"]}', headers=auth)
    assert original.status_code == 200
    assert len(original.content) == item["data"]["photos"][0]["bytes"]


def test_preview_preserves_exif_orientation(client, database, auth):
    source = Image.new("RGB", (20, 10), (70, 120, 50))
    exif = source.getexif()
    exif[274] = 6
    content = io.BytesIO()
    source.save(content, "JPEG", exif=exif)
    item = update(content.getvalue())
    assert push(client, auth, item).status_code == 200
    photo = item["data"]["photos"][0]
    assert client.put(f'/photos/{photo["id"]}', content=content.getvalue(),
                      headers={**auth, "X-Content-SHA256": photo["sha256"]}).status_code == 200
    owner = database.scalar(select(User).where(User.username == "alice"))
    login(client)
    response = client.get(f'/preview/{owner.id}/{photo["id"]}')
    assert response.status_code == 200
    assert Image.open(io.BytesIO(response.content)).size == (10, 20)


def test_web_owner_isolation_and_expert_review_version(client, database, auth, other_auth):
    item = sample_record(client, auth)
    owner = database.scalar(select(User).where(User.username == "alice"))
    path = f'/inspections/{owner.id}/{item["id"]}'
    login(client, "bob", "bob-test-password")
    assert "Осмотров пока нет" in client.get("/").text
    assert client.get(path).status_code == 404
    assert client.get(f'/preview/{owner.id}/{item["data"]["photos"][0]["id"]}').status_code == 404
    login(client)
    token = csrf(client.get(path).text)
    assert client.post(path + "/review", data={"token": token, "version": 1, "status": "confirmed"}).status_code == 403
    expert = User(username="expert", password_hash=passwords.hash("expert-test-password"), is_expert=True)
    database.add(expert)
    database.flush()
    login(client, "expert", "expert-test-password")
    token = csrf(client.get(path).text)
    assert client.post(path + "/review", data={"token": "bad", "version": 1, "status": "confirmed"}).status_code == 403
    reviewed = client.post(path + "/review", data={"token": token, "version": 1, "status": "confirmed"}, follow_redirects=False)
    assert reviewed.status_code == 303
    event = database.scalar(select(ReviewEvent))
    assert event.reviewer_id == expert.id and event.evidence_version == 1
    assert database.scalar(select(Inspection)).expert_status == "confirmed"
    assert client.post(path + "/review", data={"token": token, "version": 1, "status": "error"}).status_code == 409
    assert database.scalar(select(func.count()).select_from(ReviewEvent)) == 1
    assert "expert</small>" in client.get(path).text
    pull = client.get("/sync/inspections?after=1", headers=auth).json()
    assert pull["inspections"][0]["expert_status"] == "confirmed"


def test_xlsx_filters_and_formula_safety(client, database, auth):
    sample_record(client, auth, note="=HYPERLINK(\"bad\")\x01")
    login(client)
    reply = client.get("/export.xlsx?crop=potato")
    assert reply.status_code == 200
    book = load_workbook(io.BytesIO(reply.content))
    assert book.active.max_row == 2
    assert book.active.cell(2, 7).value.startswith("'=")
    assert book.active.cell(2, 7).data_type == "s"
    empty = load_workbook(io.BytesIO(client.get("/export.xlsx?crop=tomato").content))
    assert empty.active.max_row == 1


def test_rendered_preview_fixture(client, database, auth):
    sample_record(client, auth, note="Тестовая запись для проверки интерфейса. Не реальный диагноз.")
    login(client)
    page = client.get("/")
    assert page.status_code == 200
    # Isolated fixture, no user data; visual preview without credentials or browser automation login.
    output = Path("/results")
    if output.is_dir():
        html = page.text.replace('"/static/', '"http://127.0.0.1:8027/static/')
        html = html.replace("<main>", "<main><p>Проверка интерфейса: только синтетические тестовые записи.</p>")
        (output / "dashboard-preview.html").write_text(html, encoding="utf-8")
